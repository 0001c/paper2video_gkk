"""论文输入摄取 —— 把使用者给的文件变成流水线能用的东西。

为什么需要这一层
----------------
流水线入口 `--paper_latex_root` 要的是一个**含 main.tex 的目录**，
而使用者手里是 Word 或 PDF（2026-09-18 确认）。这个落差必须有人填。

本模块**只做提取，不做结构化**，这是刻意的分层：

    .docx/.doc/.pdf --[本模块·零成本]--> 正文文本(+标题层级)
                                    --[LLM·要花钱]--> main.tex

理由：结构化那一步需要 LLM 做章节划分，要花钱且慢；而提取是纯本地的。
拆开之后，提取质量可以单独测、反复调，不必每次烧 token。

支持的输入与路径
----------------
| 输入        | 方式                              | 质量 |
|-------------|-----------------------------------|------|
| .tex / .zip | 直接落盘（校验 main.tex 存在）      | 无损 |
| .docx       | python-docx 逐段读，保留标题层级    | 高   |
| .doc        | olefile 读 OLE2 的 WordDocument 流 | 中（有损）|
| .pdf        | PyMuPDF 提文本 + **行合并**         | 中高 |

⚠️ 为什么不用 Office / WPS 的 COM 自动化
----------------------------------------
实测（2026-09-18）`win32com` 调 `Word.Application` 时，会**连到使用者正在
运行的 WPS 实例**并挂死 2.5 分钟（`Documents.Open` 弹对话框、无返回）。
H5 后端要无人值守，不能依赖桌面应用的状态。故一律纯 Python。

⚠️ 为什么 PDF 必须做行合并
--------------------------
PDF 里的一行句子会被列宽**硬切成多行**，实测拿到过
``...conveying the researcher's scholarly id`` 这种断在半句的情况。
直接逐行拼接喂给 LLM，重建出的章节结构会跟着碎。
"""
import argparse
import hashlib
import io
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

# ── 格式嗅探：扩展名不可信，一律看 magic bytes ──
_MAGIC = (
    (b"\xd0\xcf\x11\xe0", "doc"),      # OLE2（老 .doc / .xls / .ppt）
    (b"PK\x03\x04", "zip"),            # ZIP 容器（.docx / .zip / 部分 .tex 包）
    (b"%PDF", "pdf"),
    (b"{\\rt", "rtf"),
)


def sniff_format(path):
    """返回真实格式：doc / docx / pdf / zip / tex / rtf / unknown。

    为什么不能信扩展名：使用者会把 .docx 改名成 .pdf 提交，也会拿 .doc
    当 .docx 传。按扩展名分派会在解析到一半才失败，报错点离原因很远。
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    with open(path, "rb") as handle:
        head = handle.read(8)

    kind = next((name for sig, name in _MAGIC if head.startswith(sig)), None)
    if kind == "zip":
        # ZIP 容器要再分：是 LaTeX 工程包，还是 .docx？
        try:
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
            if any(n == "word/document.xml" for n in names):
                return "docx"
            if any(n.endswith(".tex") for n in names):
                return "zip"          # LaTeX 工程包
            return "zip"
        except zipfile.BadZipFile:
            return "unknown"
    if kind:
        return kind
    # .tex 没有固定 magic，只能看内容特征。
    # ⚠️ 不能只判断「是否以 \ 开头」：LaTeX 文件常以 % 注释开头，
    # 本项目 assets/papers/gkk_mengxin/main.tex 的第一行就是中文注释，
    # 按「以 \ 开头」判会漏判成 unknown。
    with open(path, "rb") as handle:
        chunk = handle.read(8192)
    if Path(path).suffix.lower() in ('.txt', '.md'):
        try:
            text = Path(path).read_text(encoding='utf-8-sig')
        except (UnicodeError, OSError):
            return 'unknown'
        if text.strip() and not any(ord(c) < 32 and c not in '\n\r\t' for c in text):
            return 'text'
        return 'unknown'
    if any(marker in chunk for marker in
           (b"\\documentclass", b"\\begin{document}", b"\\section")):
        return "tex"
    return "unknown"


def _normalize(text):
    """统一换行、压掉多余空白与孤立控制字符。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u3000", " ")          # 全角空格
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ──────────────────────────────────────────────────────────────
# PDF：提取 + 行合并
# ──────────────────────────────────────────────────────────────
_CN_END = "。！？；：”\"’')）】》"
# 英文句号也是句末，但可能是缩写 —— 这些不算
_ABBREV = ("et al.", "e.g.", "i.e.", "fig.", "tab.", "no.",
           "vol.", "pp.", "etc.", "cf.", "vs.")
_PARA_START = re.compile(
    r"^\s*(?:"
    r"[一二三四五六七八九十]+[、.]|"          # 中文序号
    r"\d+(?:\.\d+)*[、.\s]|"                  # 数字序号
    r"第[一二三四五六七八九十\d]+[章节部分]|"
    r"(?:摘要|关键词|引言|绪论|结论|结语|参考文献|致谢|附录|Abstract|References)"
    r")"
)


def _ends_sentence(text):
    """行尾是否为句末 —— 中英文都要认。

    为什么必须单独处理英文：只放中文标点的话，英文论文的句末永远是 False，
    于是整页被合并成超长段（实测 346 行 -> 55 段、平均 325 字符），
    段落边界与标题一起被淹没。加入英文句号后粒度回到句子级。
    但英文句号也可能是缩写或人名缩写，所以要排除，否则会把
    "et al." / "J. Smith" 从中间劈开。
    """
    text = text.rstrip()
    if not text:
        return False
    if text[-1] in _CN_END:
        return True
    if text[-1] not in ".!?":
        return False
    lowered = text.lower()
    if any(lowered.endswith(item) for item in _ABBREV):
        return False
    if re.search(r"(?<![A-Za-z])[A-Z]\.$", text):    # 人名缩写 "J."
        return False
    return True


def _merge_pdf_lines(raw_lines):
    """把 PDF 的硬换行还原成「一句一行 + 空行分段」。

    判据（按优先级）：
      1. 原文空行 → 段落边界，一定断开
      2. 命中小节标记（"2.1 xxx" / "第三章" / "参考文献"）→ 断开
      3. 上一行以句末标点收尾（中英文都认）→ 断开
      4. 上一行很短（< 20 字符）→ 多半是标题，断开
      5. 其余 → 拼接（这才是被列宽切断的同一句话）

    为什么做到句子级而不是段落级：PDF 的段落边界常常没有显式空行
    （紧密排版时尤其如此），只靠空行判段落会把相邻段落粘成一团；
    而句子边界有可靠标点。按句子切开后每个语义单元都是完整的，
    下游 LLM 能据此推断段落与章节结构。
    """
    merged, buf = [], ""
    for line in raw_lines:
        stripped = line.strip()
        if not stripped:
            if buf:
                merged.append(buf)
                buf = ""
            continue

        if buf and (_PARA_START.match(stripped) or _ends_sentence(buf)
                    or len(buf) < 20):
            merged.append(buf)
            buf = stripped
        elif buf:
            # 中英混排的接缝：中文之间直接接，英文之间补空格
            joiner = "" if re.search(r"[\u4e00-\u9fff]$", buf) else " "
            buf = buf.rstrip() + joiner + stripped
        else:
            buf = stripped
    if buf:
        merged.append(buf)
    return merged


def _geometric_lines(page):
    """按基线/横坐标重组碎片字符；只去掉同位置的重复绘制。"""
    chars = [char for block in page.get_text('rawdict')['blocks'] if block['type'] == 0
             for line in block['lines'] for span in line['spans'] for char in span['chars']]
    rows = []
    for char in sorted(chars, key=lambda item: (item['origin'][1], item['origin'][0])):
        if not rows or abs(char['origin'][1] - rows[-1][0]) > 0.75:
            rows.append((char['origin'][1], []))
        rows[-1][1].append(char)
    lines, duplicates = [], 0
    for baseline, row in rows:
        kept = []
        for char in sorted(row, key=lambda item: item['origin'][0]):
            if any(char['c'] == old['c'] and abs(char['origin'][0] - old['origin'][0]) < 0.8
                   for old in kept[-4:]):
                duplicates += 1
                continue
            kept.append(char)
        lines.append(''.join(char['c'] for char in kept))
    return lines, duplicates


def has_unmapped_glyphs(text):
    return ('【未识别字形】' in text or '\ufffd' in text
            or any(ord(char) < 32 and char not in '\n\r\t' for char in text))


def extract_pdf(path):
    """PyMuPDF 提文本，逐页做行合并。返回 (正文, 元数据)。"""
    import pymupdf

    doc = pymupdf.open(path)
    pages, empty_pages, quality = [], [], []
    for idx in range(doc.page_count):
        page = doc[idx]
        raw = page.get_text()
        blocks = [block for block in page.get_text('blocks') if block[6] == 0]
        fragmented = len(blocks) > 80 and sum(len(block[4].strip()) <= 5 for block in blocks) > len(blocks) * 0.5
        duplicates = 0
        if fragmented:
            raw_lines, duplicates = _geometric_lines(page)
        else:
            raw_lines = raw.splitlines()
        controls = sum(ord(c) < 32 and c not in '\n\r\t' for c in raw)
        reasons = []
        if controls or '\ufffd' in raw:
            reasons.append('字形映射缺失（含公式），需要对照源页转录')
        if len(raw.strip()) < 20:
            reasons.append('正文过少，可能为扫描或图片页')
        # 缺失字形不能被 _normalize 静默删掉后当作完整公式。
        raw_lines = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffd]', '【未识别字形】', line) for line in raw_lines]
        lines = _merge_pdf_lines(raw_lines)
        body = "\n\n".join(lines)
        if len(body.strip()) < 20:
            empty_pages.append(idx + 1)
        pages.append('[源文第 {} 页]\n\n{}'.format(idx + 1, body))
        quality.append({'page': idx + 1, 'fragmented': fragmented, 'reading_order_repaired': fragmented,
                        'duplicate_glyphs_removed': duplicates, 'unmapped_glyphs': controls,
                        'needs_review': bool(reasons), 'reasons': reasons})
    total_chars = sum(len(p) for p in pages)
    doc.close()

    meta = {
        "format": "pdf",
        "pages": len(pages),
        "chars": total_chars,
        "empty_pages": empty_pages,      # 空页多 = 很可能是扫描版
        "likely_scanned": bool(empty_pages) and len(empty_pages) >= max(1, len(pages) // 3),
        "page_quality": quality,
        "review_pages": [item['page'] for item in quality if item['needs_review']],
        "source_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
    }
    text = "\n\n".join(pages)
    return _normalize(text), meta


# ──────────────────────────────────────────────────────────────
# DOCX：逐段读，保留标题层级
# ──────────────────────────────────────────────────────────────
def extract_docx(path):
    """python-docx 逐段读。把 Word 的 Heading 样式翻成 Markdown 标题，
    这样 LLM 判断章节结构的准确率会明显高于一堆无标记段落。
    """
    import docx

    document = docx.Document(path)
    blocks, n_head, n_para = [], 0, 0
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        match = re.match(r"Heading (\d+)", style)
        if match:
            level = min(int(match.group(1)), 6)
            blocks.append("#" * level + " " + text)
            n_head += 1
        elif style in ("Title", "Subtitle"):
            # python-docx 的 add_heading(level=0) 产生的是 'Title' 而非
            # 'Heading N'。漏掉它会让论文主标题被当成普通正文，
            # 后续 LLM 判断"这是标题页"时就少了最直接的线索。
            blocks.append("# " + text)
            n_head += 1
        else:
            blocks.append(text)
            n_para += 1

    n_images = len(document.inline_shapes)
    meta = {
        "format": "docx",
        "headings": n_head,          # 0 表示这篇 Word 没设标题样式，全平铺
        "paragraphs": n_para,
        "inline_images": n_images,
        "hint": ("未检测到任何标题样式，章节结构需靠 LLM 从正文推断"
                 if n_head == 0 else "已按 Word 的标题样式保留层级"),
    }
    return _normalize("\n\n".join(blocks)), meta


# ──────────────────────────────────────────────────────────────
# DOC（OLE2 老格式）：无 Office 依赖的粗提取
# ──────────────────────────────────────────────────────────────
# 从 .doc 的二进制流里捞正文用的字符集。
#
# ⚠️ 这里踩过两次，方向相反的坑，都记在这：
#
# ① 白名单**漏标点** → 句子被无声截断。
#    实测 `尼诺·法兰克` 的 `·`（U+00B7）漏在白名单外，
#    「摘要: 黑色电影自1946年被法国影评人尼诺」就断在那里，
#    而且没有任何报错 —— 只有拿同一篇论文的 PDF 版交叉比对才看得出来。
#
# ② 于是改成「放开可打印范围」→ 更糟。
#    `[\u0020-\uffff]` 会把 WordDocument 流里的二进制结构数据
#    也解码成"看起来像中文"的垃圾段（实测出现
#    「菢䛰拋♓叮첱썧쀪⡴…」这种），句子数从 84 虚增到 267，
#    章节命中从 5/5 掉到 2/5 —— 垃圾把正文淹没了。
#
# 结论：**要扩充白名单，而不是取消白名单**。放宽到"常见汉字 + 常见标点"，
# 既能接住 `·`「」“”…，又挡得住二进制垃圾。
_DOC_TEXT_RUN = re.compile(
    "["
    "\u4e00-\u9fff"      # CJK 统一表意文字（正文主体）
    "\u3000-\u303f"      # CJK 标点：、。《》「」【】
    "\uff00-\uffef"      # 全角：，。！？：；（）等
    "\u2010-\u2027"      # 连字符/破折号/引号/省略号：—‘’“”…
    "\u00b7\u00a0"       # 间隔号 · 、不换行空格
    "A-Za-z0-9"
    "\\s,.;:!?'\"()\\[\\]%\\-/&+=@#$*<>{}|~"
    "]{8,}")


def _garbage_ratio(text):
    """段内汉字有多大比例落在 GB2312 之外 —— 用来识别二进制解码残留。

    为什么需要：「中文字数 >= 4」挡不住这些残留（它们确实是汉字码位，
    例如「襜鉾髭谝究牨猨砨」），所以必须换判据。
    实测残留段的 GB2312 外字符占比 50%~100%，而真实正文一律 0%、表单字段也是 0%
    （表单字段虽不是论文正文，但确实是有效文本，不该在此处误杀）。
    """
    cjk = [ch for ch in text if "\u4e00" <= ch <= "\u9fff"]
    if not cjk:
        return 0.0
    outside = 0
    for ch in cjk:
        try:
            ch.encode("gb2312")
        except UnicodeEncodeError:
            outside += 1
    return outside / len(cjk)


def extract_doc(path):
    """从 OLE2 的 WordDocument 流粗提正文。

    python-docx **读不了 .doc**（只支持 .docx）。实测该格式的正文以
    UTF-16LE 存在 WordDocument 流里，所以可以直接扫可读文本段。

    局限（必须在界面上告知使用者）：
      · 拿不到样式层级，标题与正文混在一起
      · 页眉/页脚/批注可能与正文交错
      · 「快速保存」过的文档，文本顺序可能不等于阅读顺序
    所以这是**后备路径**，首选仍是引导使用者另存为 .docx。
    """
    import olefile

    ole = olefile.OleFileIO(path)
    streams = ["/".join(s) for s in ole.listdir()]
    raw = ole.openstream("WordDocument").read()
    ole.close()

    # 正文段落：连续的可读字符。过滤掉二进制垃圾（要求段内有足够的中文或字母）
    text16 = raw.decode("utf-16-le", errors="ignore")
    runs = [s.strip() for s in _DOC_TEXT_RUN.findall(text16)]
    runs = [s for s in runs
            if (len(re.findall(r"[\u4e00-\u9fff]", s)) >= 4
                or len(re.findall(r"[A-Za-z]", s)) >= 8)
            and _garbage_ratio(s) < 0.5]     # 剔除二进制解码残留

    meta = {
        "format": "doc",
        "streams": streams,
        "utf16_runs": len(runs),
        "cjk_chars": sum(len(re.findall(r"[\u4e00-\u9fff]", s)) for s in runs),
        "lossy": True,
        "hint": "老版 .doc 为有损提取：无标题层级、可能混入页眉页脚，请在预览中核对",
    }
    return _normalize("\n\n".join(runs)), meta


# ──────────────────────────────────────────────────────────────
# TEX / ZIP：直接落盘，不做任何文本处理
# ──────────────────────────────────────────────────────────────
def extract_latex(path, out_dir):
    """.tex 或 LaTeX 工程包：解压/复制到 out_dir，并确认 main.tex 存在。"""
    target = os.path.join(out_dir, "latex_proj")
    if os.path.isdir(target):
        shutil.rmtree(target)
    os.makedirs(target, exist_ok=True)

    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as zf:
            zf.extractall(target)
        # 工程包常见再套一层目录，往下找 main.tex
        entries = [e for e in os.listdir(target)]
        if len(entries) == 1 and os.path.isdir(os.path.join(target, entries[0])):
            inner = os.path.join(target, entries[0])
            if any(f.endswith(".tex") for f in os.listdir(inner)):
                target = inner
    else:
        shutil.copy2(path, os.path.join(target, "main.tex"))

    mains = [f for f in os.listdir(target) if f == "main.tex"]
    tex_files = [f for f in os.listdir(target) if f.endswith(".tex")]
    if not mains:
        # 没有 main.tex 但有别的 .tex：帮它改名为 main.tex（流水线认这个名字）
        if len(tex_files) == 1:
            os.rename(os.path.join(target, tex_files[0]),
                      os.path.join(target, "main.tex"))
            mains = ["main.tex"]
    meta = {
        "format": "latex",
        "latex_root": os.path.abspath(target),
        "has_main_tex": bool(mains),
        "tex_files": tex_files,
        "needs_llm": False,
    }
    if not mains:
        raise RuntimeError(
            "LaTeX 输入里找不到 main.tex（也没有唯一的 .tex 可改名）。"
            "请在压缩包里确保入口文件名为 main.tex。目录内容：{}".format(
                sorted(os.listdir(target))[:20]))
    return "", meta


# ──────────────────────────────────────────────────────────────
# 统一入口
# ──────────────────────────────────────────────────────────────
# ⚠️ **改了本模块的提取逻辑就要把这个号 +1**。
# 调用方（web/app.py 的内容库）用它做缓存键的一部分：
# 同一个文件 + 同一个版本号才允许直接复用已提取的正文。
# 不 +1 的后果是"改了预处理却仍命中旧缓存"——内容对不上，且没有任何报错。
EXTRACTOR_VERSION = "2"

_EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "doc": extract_doc,
}


def ingest(src, out_dir):
    """把任意支持的输入变成「正文文本 + 元数据」。

    返回 ``(text, meta)``；``meta['needs_llm']`` 为 False 时表示无需
    结构化（LaTeX 直接可用）。
    """
    os.makedirs(out_dir, exist_ok=True)
    fmt = sniff_format(src)
    print("[ingest] {} -> 嗅探为 {}".format(os.path.basename(src), fmt))

    if fmt in ("tex", "zip"):
        text, meta = extract_latex(src, out_dir)
    elif fmt in _EXTRACTORS:
        text, meta = _EXTRACTORS[fmt](src)
    elif fmt == 'text':
        text = Path(src).read_text(encoding='utf-8-sig').strip()
        meta = {'format': 'text', 'chars': len(text), 'needs_llm': True}
    elif fmt == "rtf":
        raise RuntimeError(
            "检测到 RTF 格式。请在 Word / WPS 里另存为 .docx 后再上传。")
    else:
        raise RuntimeError(
            "无法识别的文件格式（前 8 字节既不是 OLE2 / ZIP / PDF，也不是 LaTeX）。"
            "支持：.tex / .zip / .docx / .doc / .pdf / UTF-8 .txt")

    meta.setdefault("needs_llm", True)
    meta["source"] = os.path.abspath(src)
    meta["detected_format"] = fmt
    if fmt == 'pdf' and meta.get('review_pages'):
        import pymupdf
        evidence_dir = os.path.join(out_dir, 'source_pages')
        os.makedirs(evidence_dir, exist_ok=True)
        with pymupdf.open(src) as document:
            for item in meta['page_quality']:
                if item['needs_review']:
                    filename = os.path.join(evidence_dir, 'page-{:04d}.png'.format(item['page']))
                    document[item['page'] - 1].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(filename)
                    item['source_image'] = os.path.abspath(filename)

    if text:
        text_path = os.path.join(out_dir, "paper_text.md")
        with io.open(text_path, "w", encoding="utf-8", newline="") as handle:
            handle.write(text)
        meta["text_path"] = text_path
        meta["chars"] = len(text)
    with io.open(os.path.join(out_dir, "ingest_meta.json"), "w",
                 encoding="utf-8", newline="") as handle:
        json.dump(meta, handle, ensure_ascii=False, indent=2)

    return text, meta


def main(argv=None):
    parser = argparse.ArgumentParser(description="论文输入摄取（tex/zip/docx/doc/pdf）")
    parser.add_argument("source", help="输入文件路径")
    parser.add_argument("--out", default="./.ingest_out", help="输出目录")
    parser.add_argument("--preview", type=int, default=700,
                        help="打印正文前 N 个字符")
    args = parser.parse_args(argv)

    text, meta = ingest(args.source, args.out)
    print()
    print("=== 元数据 ===")
    for key, value in meta.items():
        if key in ("source",):
            continue
        print("   %-16s %s" % (key, value))
    if text:
        print()
        print("=== 正文预览（前 %d 字符）===" % args.preview)
        print(text[:args.preview])
    return 0


if __name__ == "__main__":
    sys.exit(main())
