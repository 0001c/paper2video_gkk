"""幻灯片排错回归：图片尺寸、纠错约束、Word 图注；无真实模型请求。"""
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'Paper2Video-main/src')]
import pv_env  # noqa: F401
import paper_figures
import slide_quality
import slide_code_gen_select_improvement as slides
from docx import Document
from PIL import Image


class LayoutChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / 'jobs/_validation')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_images_get_height_bound_without_changing_content(self):
        code = r'\item 真实性提升50%。\includegraphics[width=\textwidth]{figures/a.png}'
        bounded = slide_quality.bound_image_sizes(code)
        self.assertIn(r'width=0.9\linewidth,height=0.45\textheight,keepaspectratio', bounded)
        self.assertIn(r'\item 真实性提升50%。', bounded)
        self.assertIn('{figures/a.png}', bounded)
        self.assertNotIn('tiny', bounded)
        self.assertEqual(bounded, slide_quality.bound_image_sizes(bounded))

    def test_frame_count_is_a_quality_gate(self):
        code = r'\begin{frame}first\end{frame}\begin{frame}second\end{frame}'
        self.assertEqual(slide_quality.quality_errors(code, expected_count=2), [])
        errors = slide_quality.quality_errors(code, expected_count=1)
        self.assertEqual(len(errors), 1)
        self.assertIn('要求1页，实际2帧', errors[0])

    def test_rowspec_repair_preserves_options_and_table_content(self):
        code = ('\\begin{tblr}{colspec={Q[l]Q[r]},\n'
                ' rowspec = { Q[m,bg=red] ,\n Q[t,fg=blue], Q[b] },\n'
                ' row{1}={font=\\bfseries,fg=red}}\n'
                'A, B & 1,234 \\\\\n\\end{tblr}')
        expected = code.replace('] ,\n Q', '] \n Q').replace('], Q', '] Q')
        repaired = slide_quality.normalize_rowspec(code)
        self.assertEqual(repaired, expected)
        self.assertEqual(slide_quality.normalize_rowspec(repaired), repaired)

    def test_rowspec_repair_leaves_valid_and_unknown_syntax_unchanged(self):
        for code in (
            'rowspec={Q[m]Q[m]},',
            'rowspec={Q[m,bg=red] Q[t,fg=blue]},',
            'colspec={Q[m], Q[m]}, rows={m},',
            '% rowspec={Q[m], Q[m]},\n',
            '\\\\% rowspec={Q[m], Q[m]},\n',
            r'rowspec={Q[m,font={\bfseries}], Q[t]},',
            r'rowspec={Q[m], \customrow},',
        ):
            with self.subTest(code=code):
                self.assertEqual(slide_quality.normalize_rowspec(code), code)
        self.assertEqual(slide_quality.normalize_rowspec(r'\% rowspec={Q[m], Q[m]},'),
                         r'\% rowspec={Q[m] Q[m]},')

    def test_compile_repairs_rowspec_before_running_compiler(self):
        target = self.root / 'slides.tex'
        code = '\\documentclass{article}\nrowspec={Q[m], Q[m]},'
        target.write_text(code, encoding='utf-8')

        def compile_process(*args, **kwargs):
            self.assertEqual(target.read_text(encoding='utf-8'), code.replace('], Q', '] Q'))
            return SimpleNamespace(returncode=0, stdout='', stderr='')

        with patch.object(slides.subprocess, 'run', side_effect=compile_process):
            self.assertTrue(slides.compile_tex(target)[0])

    def source_frame(self, caption='图1 概念模型', fid='F1'):
        return ('\\begin{frame}{研究方法}\n% source-figure: ' + fid + '\n'
                '\\begin{figure}\n\\includegraphics{figures/research.png}\n'
                '\\caption{\\tiny ' + caption + '}\n\\end{figure}\n\\end{frame}')

    def test_adjacent_source_marker_is_moved_without_guessing(self):
        frame = self.source_frame()
        normalized = slide_quality.normalize_source_bindings(frame)
        self.assertIn('\\begin{figure}\n% source-figure: F1', normalized)
        self.assertEqual(slide_quality.normalize_source_bindings(normalized), normalized)
        figure = {'rel': 'figures/research.png', 'figure_id': 'F1', 'caption': '图1 概念模型'}
        self.assertEqual(slide_quality.quality_errors(frame, figures=[figure]), [])
        self.assertTrue(slide_quality.quality_errors(frame.replace('F1', 'F2'), figures=[figure]))
        self.assertTrue(slide_quality.quality_errors(frame.replace('% source-figure: F1\n', ''), figures=[figure]))
        distant = frame.replace('F1\n', 'F1\n\\item 不同内容\n')
        self.assertEqual(slide_quality.normalize_source_bindings(distant), distant)
        self.assertTrue(slide_quality.quality_errors(distant, figures=[figure]))

    def test_composite_source_preserves_all_original_numbers(self):
        caption = '图5 内容真实性灵敏度分析 图6 内容准确性灵敏度分析'
        figure = {'rel': 'figures/research.png', 'figure_id': 'F5', 'caption': caption}
        frame = self.source_frame(caption.replace(' 图6', '; 图6'), 'F5')
        self.assertEqual(slide_quality.quality_errors(frame, figures=[figure]), [])
        for bad in (frame.replace('图6', '图7'), frame.replace('; 图6 内容准确性灵敏度分析', ''),
                    frame.replace('内容真实性', '制度规范'), frame.replace('F5', 'F6'),
                    frame.replace('F5\n', 'F5\n% source-figure: F6\n')):
            with self.subTest(code=bad):
                self.assertTrue(slide_quality.quality_errors(bad, figures=[figure]))

    def test_caption_finder_does_not_bind_body_mentions(self):
        body = (0, 101, 100, 111, '响较为显著，如图9 和图10 所示。', 0, 0)
        caption = (0, 75, 100, 85, '图5 内容真实性 图6 内容准确性', 0, 0)
        self.assertIsNone(paper_figures.figure_id(body[4]))
        self.assertEqual(paper_figures.figure_id(' Figure 2-15 Research network'), 'F2-15')
        self.assertEqual(paper_figures._caption_for(None, (0, 90, 100, 100), [body, caption]), caption[4])
        self.assertEqual(paper_figures._caption_for(None, (0, 90, 100, 100), [body]), '')

    def test_legacy_quality_report_is_rechecked(self):
        import json
        target = self.root / 'slides.tex'
        target.write_text(r'\documentclass{beamer}', encoding='utf-8')
        target.with_suffix('.pdf').write_bytes(b'fixture')
        slide_quality.record_quality(target, '', True)
        report_path = target.with_suffix('.quality.json')
        report = json.loads(report_path.read_text(encoding='utf-8'))
        report['version'] = 4
        report_path.write_text(json.dumps(report), encoding='utf-8')
        with patch.object(slides, 'compile_tex', return_value=(True, '')) as compiler:
            slide_quality.require_quality(self.root)
            compiler.assert_called_once_with(target)

    def test_rendered_page_count_detects_overlays(self):
        import pymupdf
        target = self.root / 'slides.tex'
        target.write_text(r'\documentclass{beamer}\begin{document}\begin{frame}one frame\end{frame}\end{document}', encoding='utf-8')
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.new_page()
            pdf.save(target.with_suffix('.pdf'))
        errors = slide_quality.record_quality(target, '', True, expected_count=1)
        self.assertTrue(any('渲染PDF实际2页' in error for error in errors))

    def test_existing_small_image_and_special_size_are_preserved(self):
        code = r'\includegraphics[width=.3\linewidth,height=.2\textheight,keepaspectratio=false]{a.png}'
        bounded = slide_quality.bound_image_sizes(code)
        self.assertIn(r'width=0.3\linewidth,height=0.2\textheight,keepaspectratio', bounded)
        self.assertNotIn('false', bounded)
        special = r'\includegraphics[width=3cm,height=2cm]{b.png}'
        self.assertIn('width=3cm,height=2cm', slide_quality.bound_image_sizes(special))

    def test_unchanged_correction_stops_and_keeps_count_and_figure_constraints(self):
        code = r'\documentclass{beamer}\begin{document}\begin{frame}content\end{frame}\end{document}'
        prompt = self.root / 'prompt.txt'
        prompt.write_text('生成课件', encoding='utf-8')
        target = self.root / 'slides.tex'
        agent = MagicMock()
        agent.step.return_value = SimpleNamespace(msgs=[SimpleNamespace(content=code)], info={'usage': {}})
        feedback = 'warning: noise\nerror: 版面裁切: 100pt\nerror: 版面裁切: 100pt\nwarning: unrelated noise'
        constraints = '恰好20页；figure_id=F5；图5内容真实性灵敏度分析'
        with patch.object(slides, 'create_agent', return_value=agent), \
                patch.object(slides, 'compile_tex', return_value=(False, feedback)), \
                patch.object(slides, 'correct_prompt_path', str(ROOT / 'Paper2Video-main/src/prompts/slide_beamer_correct.txt')):
            with self.assertRaises(RuntimeError) as error:
                slides.latex_code_gen_upgrade(prompt, self.root, target, {}, {},
                                              if_tree_search=False, extra_prompt=constraints)
        self.assertEqual(agent.step.call_count, 2)
        correction = agent.step.call_args.args[0].content
        self.assertIn(constraints, correction)
        self.assertIn('Every \\caption must be inside', correction)
        self.assertIn('height<=0.5', correction)
        self.assertNotIn('unrelated noise', str(error.exception))
        self.assertEqual(str(error.exception).count('100pt'), 1)
        self.assertEqual(target.read_text(encoding='utf-8'), code)

    def test_first_generation_applies_local_image_bounds(self):
        code = r'\documentclass{beamer}\begin{document}\begin{frame}\includegraphics{a.png}\begin{tblr}{rowspec={Q[m], Q[m]}}A&B\\C&D\end{tblr}\end{frame}\end{document}'
        prompt = self.root / 'prompt.txt'
        prompt.write_text('生成课件', encoding='utf-8')
        target = self.root / 'slides.tex'
        agent = MagicMock()
        agent.step.return_value = SimpleNamespace(msgs=[SimpleNamespace(content=code)], info={'usage': {}})
        with patch.object(slides, 'create_agent', return_value=agent), patch.object(slides, 'compile_tex', return_value=(True, '')):
            slides.latex_code_gen_upgrade(prompt, self.root, str(target), {}, {}, if_tree_search=False)
        self.assertEqual(agent.step.call_count, 1)
        self.assertIn(r'height=0.45\textheight', target.read_text(encoding='utf-8'))
        self.assertIn('rowspec={Q[m] Q[m]}', target.read_text(encoding='utf-8'))

    def image(self, name, color):
        path = self.root / name
        Image.new('RGB', (220, 200), color).save(path)
        return str(path)

    def test_docx_source_order_paired_captions_and_decoration_exclusion(self):
        document = Document()
        document.add_picture(self.image('cover.png', 'red'))
        document.add_paragraph('期刊封面说明')
        document.add_picture(self.image('research.png', 'blue'))
        for _ in range(15):
            document.add_paragraph('')
        document.add_paragraph('图 10 第一张真实研究图')
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(self.image('left.png', 'green'))
        paragraph.add_run().add_picture(self.image('right.png', 'yellow'))
        document.add_paragraph('图 2 左图    图 3 右图')
        source = self.root / 'source.docx'
        document.save(source)
        figures, _ = paper_figures.extract(str(source), str(self.root / 'project'), log=lambda _: None)
        self.assertEqual([f['figure_id'] for f in figures], ['F10', 'F2', 'F3'])
        self.assertEqual([f['caption'] for f in figures], ['图 10 第一张真实研究图', '图 2 左图', '图 3 右图'])
        with Image.open(self.root / 'project' / figures[0]['rel']) as image:
            self.assertEqual(image.getpixel((0, 0)), (0, 0, 255))
        self.assertTrue(all(f['docx_media'].startswith('word/media/') for f in figures))

    def test_unlabeled_docx_does_not_guess_caption(self):
        document = Document()
        document.add_picture(self.image('unlabeled.png', 'blue'))
        document.add_paragraph('普通文字，无图注。')
        source = self.root / 'unlabeled.docx'
        document.save(source)
        figures, _ = paper_figures.extract(str(source), str(self.root / 'project'), log=lambda _: None)
        self.assertEqual(len(figures), 1)
        self.assertIsNone(figures[0]['figure_id'])
        self.assertEqual(figures[0]['caption'], '')

    def test_caption_cannot_be_stolen_by_next_unlabeled_image(self):
        document = Document()
        document.add_picture(self.image('first.png', 'blue'))
        document.add_paragraph('图 1 已绑定研究图')
        document.add_picture(self.image('decoration.png', 'red'))
        document.add_paragraph('普通正文')
        source = self.root / 'ambiguous.docx'
        document.save(source)
        figures, _ = paper_figures.extract(str(source), str(self.root / 'project'), log=lambda _: None)
        self.assertEqual([f['figure_id'] for f in figures], ['F1'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
