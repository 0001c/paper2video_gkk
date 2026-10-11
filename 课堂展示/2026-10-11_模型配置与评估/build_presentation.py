"""离线生成可编辑 PPT、逐页讲稿、PDF 与预览；不调用项目模型或网络。"""
from pathlib import Path
import hashlib
import json
import re
from io import BytesIO
import shutil

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as DocPt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
W, H = 1600, 900
FONT = Path('C:/Windows/Fonts/msyh.ttc')
BOLD = Path('C:/Windows/Fonts/msyhbd.ttc')
BG, INK, GREEN, MUTED = '#F5F3EB', '#173A32', '#2D6B56', '#536860'
PALE, LINE, WHITE = '#E5EDE4', '#CCD7CD', '#FFFFFF'
AMBER, RED = '#936817', '#AA4439'
NAME = 'Paper2Video_翻转课堂_模型配置与评估_15分钟'

SOURCES = {
 'P': ('论文', 'docs/Paper2Video.pdf', 'arXiv:2510.05096v2，2025-10-09；19页'),
 'C': ('模型配置', 'web/model_config.py', '角色、参数范围、执行能力校验'),
 'F': ('前端默认', 'web/frontend/src/modelConfig.ts', '预设型号、新任务评价默认值'),
 'R': ('模型运行', 'Paper2Video-main/src/model_runtime.py', '角色客户端、任务快照、禁止全局凭据回落'),
 'Q': ('问答实现', 'Paper2Video-main/src/qa_evaluation.py', '原文题库、证据、抽帧、评分与缓存'),
 'D': ('问答说明', 'docs/双阶段问答评价与数字人成本控制.md', '双阶段门槛、人工决策、独立任务'),
 'H': ('历史实测', 'docs/自然语言处理第二章全链路实测评估.md', '2026-10-08历史样本，不代表当前新版'),
 'L': ('修复验收', 'docs/全链路P1与P2修复与离线验收.md', '字幕区、例题、逐项门禁及离线验收'),
 'V': ('最新验收', 'log/2026-10-11_17-30-00_数字人成本提醒与双阶段问答评价.md', '离线回归与实际页面检查；未实测付费问答'),
}

# 16页主讲900秒；末两页为备用，不计入15分钟。
SLIDES = [
 dict(title='把论文变成可审核的教学视频', section='开场', seconds=35, sources=['P','D'],
      note='大家好，今天我展示的是 Paper2Video 项目，以及我们围绕教学场景做的改造。先请大家想象一个问题：把一章教材自动变成视频，文件能够播放，就能拿去教学了吗？我们的答案是，还需要检查讲了什么、声音是否完整、画面能否支持理解。接下来十五分钟，我重点讲两件事：第一，文字、视觉、语音等模型怎样分工和配置；第二，如何借鉴论文的评价体系，把评价放进生成流程。希望大家最后能判断，一段教学视频什么时候可以继续生成，什么时候必须停下来复核。', cue='开场提问停顿3秒；点出“配置”和“评价”两条主线。'),
 dict(title='论文解决什么问题？课堂需要什么？', section='论文与场景', seconds=45, sources=['P'],
      note='先明确三个名字。Paper2Video 是论文提出的任务与基准；PaperTalker 是论文中的生成框架；我们当前项目是在其代码基础上，面向中文教材和课堂展示做工程改造。论文发表于二〇二五年十月的这个预印本版本，收集了一百零一组研究论文及作者演示视频，视频平均六分十五秒。它关心的是长文档中的文字、公式、图表怎样转成彼此对齐的课件、声音、字幕和演讲者。我们进一步把使用场景放进翻转课堂：课前学生观看，课中教师讨论错题和例题，课后继续修订材料。这里教师审核是教学流程的一部分。', cue='解释Paper2Video与PaperTalker不是两个模型。'),
 dict(title='四个 Builder，串起多模态生成', section='方法框架', seconds=50, sources=['P','R'],
      note='论文把生成拆成四个 Builder。Slide Builder 写 Beamer 课件并根据编译反馈修复；Subtitle Builder 读取课件画面，生成讲解和视觉焦点；Cursor Builder 把焦点变成画面位置，并与语音时间对齐；Talker Builder 负责配音和演讲者视频。论文中的 Tree Search Visual Choice 会先生成不同字号或图像尺寸的候选，再让视觉模型选择更好的布局。当前项目保留这些分工，加入源文核对、讲稿审核、费用和双阶段问答。布局搜索在网页任务中是可配置项，默认关闭。论文按页并行生成数字人，而本项目的本地任务使用串行队列，讲稿请求可以单独设置并发。', cue='沿流程从左到右指示；强调论文设计与本机实现的区别。'),
 dict(title='模型按角色选择，型号可以替换', section='模型配置', seconds=70, sources=['P','C','F'],
      note='这里是模型配置的核心。论文实验使用 GPT-4.1 作为视觉语言模型、Gemini-2.5-Flash 作为视频评价模型；语音使用 F5-TTS，光标定位使用 UI-TARS，时间对齐使用 WhisperX，演讲者使用 Hallo2 或 FantasyTalking。当前项目并没有照搬全部型号。网页百炼预设中，文字和课件是 qwen3.8-flash，视觉讲稿和定位使用 qwen3-vl-plus，配音使用 CosyVoice-v3.5-plus；成片识别保留本地 Qwen3-ASR。数字人可选本地 MuseTalk 或云端 wan2.2-s2v，新任务默认关闭。问答评价是独立角色，可以与生成模型采用不同连接。这里列的是代码预设，不是所有历史实验都用了这些型号。关键是为每个角色选择合适能力，而不是找一个最贵的模型包办所有环节。', cue='表格从任务读到模型；ASR与数字人两行重点说明。'),
 dict(title='配置要能复现，也要能在执行前校验', section='模型配置', seconds=65, sources=['C','F','R'],
      note='一个模型名字还不够。每个角色至少要记录运行模式、提供方、连接引用、型号和请求参数。文字、视觉和问答角色可以调整 temperature 和最大输出 token；留空意味着沿用模型接口默认值，项目没有统一的最佳温度。配置层允许温度零到二，输出上限还会结合型号目录进一步校验。云端配音语速是零点五到二，本地配音目前只支持默认语速。讲稿并发默认三，可在一到八之间调整，它只控制讲稿请求，不代表本地 GPU 可以同时跑八个任务。开始执行时，系统保存无密钥的配置快照，按角色使用对应连接。连接失效就要求重新绑定。这样才能解释，为什么同一素材的两次结果不同，而不是只留下一个视频文件。', cue='示例仅展示单角色配置；参数示例不是推荐最优值。'),
 dict(title='本地 / API：按资源和环节组合', section='模型配置', seconds=55, sources=['C','F','D'],
      note='本项目支持两种预设，也允许逐角色自定义。API 路线可以减少本地生成模型的加载压力，但会上传相应材料，产生调用费用，并依赖服务连接。本地路线需要准备权重、显存和运行环境，也会消耗时间与算力。需要特别说明，本地预设不等于全链路完全离线：文字、视觉和问答仍需要明确的兼容服务，可以由部署者提供本地服务。与此同时，当前成片问答强制使用本地 Qwen3-ASR，所以百炼预设也不是纯云方案。如果云端适配器没有实现，系统会报错，而不是自动换到本地。课堂展示时，我会先用已有截图讲清配置，再演示只读结果，避免现场等待完整生成。', cue='用两列比较；讲清两种预设的实际边界。'),
 dict(title='对齐靠真实时间轴，不靠平均分配', section='模型配置', seconds=55, sources=['P','C','L'],
      note='光标和字幕需要同时解决两个问题：指向哪里，什么时候出现。论文用 UI-TARS 定位，再用 WhisperX 提取字词时间。当前项目会优先用 PDF 文字定位，找不到合适目标时，再走配置的定位模型。云配音直接使用合成服务返回的真实字词时间，本地配音则使用 ForcedAligner。口播中的数字或公式与云端转写不一致时，系统可能拒绝合片，这时应修正口播和映射，不能平均切时间来伪装成功。新版字幕沿真实字词边界拆句，增加独立字幕区域；不需要指示时隐藏光标。数字人启用后，音频按目标十八秒拆分，并严格保证每段小于二十秒，检查采样是否连续。', cue='指示“口播→时间轴→字幕/光标”；十八秒是分段目标。'),
 dict(title='配置选择，也会改变费用与重试成本', section='模型配置', seconds=60, sources=['H','D'],
      note='费用也属于配置的一部分。第二章历史实测累计标准价估算为五点六六九四元，覆盖了一百二十二条去重费用记录。这次包含视觉转录对照、失败后的配音恢复，以及一个 Wan 短片，所以不能把它当成生成一章的固定报价。账本中 Wan 检测和生成约三点〇六九元，配音约二点三七三二元；主视频并没有整章数字人。启用数字人或自动重生成，都可能继续增加费用。当前新任务默认关闭数字人，自动重生成默认最多两次，超过后交给用户决策。缓存只有在素材、模型和参数等指纹一致且通过校验时才复用。失败请求若已产生用量，也应计入累计成本。', cue='金额为历史标准价估算，非账单；不要泛化成每章价格。'),
 dict(title='论文用四种评价，回答四类问题', section='评估体系', seconds=55, sources=['P','D'],
      note='论文的评价体系很适合帮助我们拆开质量问题。Meta Similarity 比较生成内容和真人材料，并比较音色嵌入；它回答的是相似性。PresentArena 让视频模型比较两个视频，并交换输入顺序，观察偏好。PresentQuiz 从原始论文出四选一题，再让模型依据视频回答，用来衡量信息覆盖。IP Memory 则让模型根据演讲者肖像，在四组短片和问题之间做关联，关注作者辨识。它不是隔天知识记忆测试。当前项目优先集成 PresentQuiz 的思路，其他研究脚本虽然保留，却没有全部集成到网页。我们还需要单独检查可读性、事实和声音，不能把这四种指标简单平均成教学总分。', cue='四项评价各用一句话解释；标注当前集成范围。'),
 dict(title='固定原文题库，避免“自己出题自己过关”', section='问答评价', seconds=65, sources=['Q','D'],
      note='为什么出题必须来自核对过的原文？假设讲稿把某个定义写错了，如果根据这份讲稿出题，错误定义也会变成标准答案，系统就可能自我通过。当前实现从核对后的原文或选定教学范围出题，默认二十道四选一题，包含细节和理解类别。每题保存标准答案和逐字原文摘录，并检查摘录确实存在。答题请求只发送题目、选项和待评价材料，不发送答案与原文依据。两轮在条件不变时使用固定题库。答对且提供可核验材料依据，才计为有效正确题。二十题意味着一题就是五个百分点，门槛附近的变化应逐题检查。需要承认，题目和标准答案的语义仍可能有问题，人工复核不能被摘录校验替代。', cue='用数据流说明答案留在评分端；示例题是课堂示意。'),
 dict(title='讲稿评价与成片评价，使用不同证据', section='问答评价', seconds=65, sources=['Q','D'],
      note='双阶段的意义，在于评价对象发生了变化。讲稿阶段只读取实际口播文本，不把光标指令当作教学内容。成片阶段必须从最终视频独立提取完整音轨，每六十秒一段，用本地 Qwen3-ASR 识别，并均匀抽取四到二十四张实际画面。这里没有把原讲稿拿来冒充成片证据，也不是直接把完整视频交给论文中的视频模型。两次分数如果下降，可能与配音缺失、识别错误、图示不清或材料变化有关，要查看证据才能定位。文本依据检查逐字匹配；画面依据检查抽样时间，但画面文字是否真实仍依赖评委。报告保存原文、材料、模型与题库指纹，材料一旦改变，旧评价就不能继续放行。', cue='强调“音轨来自最终视频”；抽帧有上限，无法覆盖每个画面。'),
 dict(title='评分触发动作，异常回到人工复核', section='评价闭环', seconds=60, sources=['C','Q','D'],
      note='默认规则是：讲稿达到百分之八十，可以自动进入配音合片，但教学结构审核也必须通过。成片低于百分之六十，系统根据错题和原文依据调整讲稿提示词，生成新任务；六十到不足八十之间，等待用户接受或重生成；八十及以上完成。如果最多两次自动重生成都没有通过，就暂停。现在请大家快速判断：讲稿十七题正确，成片十四题正确，应该自动完成吗？停顿后给出答案：不应该，成片是百分之七十，处在人工决策区间。另一个情况是识别异常，它没有有效分数，应暂停复核，而不是记零分触发重做。这些阈值是可调整的工程策略，不是论文规定的教学合格标准。', cue='互动预留12秒；答案是人工决策，不能自动完成。'),
 dict(title='论文结果：模型答得好，真人仍更偏好人工', section='论文实验', seconds=65, sources=['P'],
      note='这页只看论文实验，不是我们的项目测试。表二中，PaperTalker 的细节问答准确率是八十四点二，理解题是九十五点一；真人视频分别是七十三点八和九十点八。看起来生成视频在模型问答上更强。但图六的人工偏好评价中，真人视频得四点六分，PaperTalker 得三点八分，评价抽取十篇论文、十名参与者。这说明信息覆盖和观看偏好测量的是不同维度。论文还报告了光标定位任务中有光标与无光标的准确率差异，但那是位置识别，不能解释成学生学习能力提高。同样，论文在八张 A6000 上测试的生成效率也不能直接套到本机。真正有用的启发是：同一视频要用不同证据，从多个维度评价。', cue='左边单位为%，右边为1–5分，不能直接相互比较。'),
 dict(title='项目实测：技术跑通以后，仍要审教学', section='项目证据', seconds=60, sources=['H','L','V'],
      note='回到本项目。二〇二六年十月八日第二章历史实测，十七页教材生成十二页课件，人工辅助恢复后，主视频时长二十五分二十四秒，全程解码通过。但是报告仍发现字幕遮挡、课件裁切和讲稿事实问题。画面中这段长字幕覆盖了网络图和图注；讲稿里还有梯度与负梯度、正向微分与数值差分混淆。这就是为什么，能播放不能证明能教学。之后代码已经加入独立字幕区、例题和结构门禁，并在十月十一日加入双阶段问答。最新版通过了十六项问答离线回归，识别和模型使用替身；还没有新的真实问答成绩，也没有学生学习实验。因此旧样本揭示问题，新版测试证明流程，二者证据不能混用。', cue='图片为历史成片抽帧；不要说新版仍有该遮挡或已经真实解决全章。'),
 dict(title='课堂演示：配置、证据、决定', section='翻转课堂互动', seconds=50, sources=['F','D','V'],
      note='课堂展示可以这样组织。课前让学生带着三个问题看短视频：模型怎样分工、评价依据从哪里来、低分怎样处理。课中先展示模型配置，再看双评分卡与逐题依据，让同学决定接受还是重生成。右侧截图中八十五与七十是离线测试夹具，用于展示界面，不能当成真实项目成绩。请大家再选一个最需要验证的因素：更换评价模型、增加字幕安全区，还是开启数字人？这些都能作为后续实验，但每次只改变一个因素，其他材料和题库固定。最后用原文复核错题，并补充学生观看前后测。今天演示已有截图就可以完成这条流程，不需要现场生成二十五分钟视频。', cue='互动15秒；可接受多种选择，但要求说明控制变量。'),
 dict(title='下一步：把自动评分接回真实学习', section='收束', seconds=45, sources=['P','D','V'],
      note='最后回到今天的目标。模型配置的价值，是让每个环节职责明确、参数可追溯、执行可校验。评价的价值，是把发现的问题转成下一步动作。论文提供了多模态分工和四类评价思路；当前项目把其中的原文问答思路扩展为讲稿与成片两轮，并加入阈值、费用控制和人工决策。接下来要做的是，选一组人工复核题，真实运行新版评价，再固定题库比较修改前后的产物，补充连续听审和学生前后测。我们能展示的是一个可以审核、可以追溯的生成系统；能否真正提高学习效果，还需要真实课堂的数据来回答。谢谢大家。', cue='用三句话收束：角色分工、证据评分、人工决策。'),
 dict(title='备用：常见提问与回答', section='答辩备用', seconds=0, sources=['P','C','Q','D'],
      note='问：为什么不只用一个大模型？答：生成、语音、定位和评价需要不同能力与接口，按角色拆分更便于替换与定位故障。问：八十的阈值从哪里来？答：当前工程默认，可配置，尚未用学生效果标定。问：题目有原文摘录就一定正确吗？答：只能证明摘录存在，题意、答案与推断仍需人工复核。问：问答高分能否说明视频没有错误？答：题库和抽帧有限，必须补源文事实审核与连续听审。问：IP Memory 是否等于学习记忆？答：论文做的是肖像与视频问题关联，不是隔天概念测验。问：本地模式完全离线吗？答：需要单独部署文字、视觉和问答兼容服务，当前预设本身不构成完全离线承诺。', cue='备用，不计入主讲15分钟。'),
 dict(title='备用：来源与复现入口', section='证据索引', seconds=0, sources=list(SOURCES),
      note='论文引用按PDF显示页码：方法第五至七页；评价第四至五页及第十四页；表二和图六在第八页，课件质量与光标消融在第九页。当前项目以代码和最新验收日志为准，早期README存在阶段描述，不能作为当前能力唯一依据。配套讲稿包含完整来源表、时间分配、演示步骤与口径说明。全部课件生成在本地进行，未创建生产任务、未调用付费模型。', cue='备用，不计入主讲15分钟。'),
]


def rgb(color):
    return RGBColor.from_string(color.lstrip('#'))


def font(size, bold=False):
    return ImageFont.truetype(str(BOLD if bold else FONT), int(size))


def wrap(text, width, size, bold=False):
    f = font(size, bold)
    result = []
    for paragraph in text.split('\n'):
        # 拉丁型号/参数名保持整体，中文按字换行。
        tokens = re.findall(r'[A-Za-z0-9_.%/+-]+|.', paragraph)
        line = ''
        for token in tokens:
            if line and f.getlength(line + token) > width:
                result.append(line)
                line = token.lstrip()
            else:
                line += token
        result.append(line)
    return result


class Page:
    def __init__(self, prs, index, spec, dark=False):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.image = Image.new('RGB', (W, H), INK if dark else BG)
        self.draw = ImageDraw.Draw(self.image)
        self.dark = dark
        self.fg = BG if dark else INK
        self.bg = INK if dark else BG
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(self.bg)
        self.index, self.spec = index, spec
        self.text(76, 45, 1440, 38, f'PAPER2VIDEO  /  {spec["section"]}', 21, '#B5D6C0' if dark else GREEN, True)
        if index != 0:
            self.text(76, 107, 1450, 86, spec['title'], 51, self.fg, True)
        self.rect(76, 812, 1448, 2, '#517265' if dark else LINE)
        refs = ' · '.join(spec['sources'])
        foot = f'论文 × 项目实践  |  2026.10.11  |  来源 {refs}'
        self.text(76, 834, 1290, 30, foot, 18, '#B5D6C0' if dark else MUTED)
        self.text(1410, 830, 110, 35, f'{index+1:02d} / 18', 21, '#B5D6C0' if dark else MUTED)

    def rect(self, x, y, w, h, fill, stroke=None):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/120), Inches(y/120), Inches(w/120), Inches(h/120))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            self.draw.rectangle((x, y, x+w, y+h), fill=fill, outline=stroke, width=2)
        else:
            shape.line.fill.background()
            self.draw.rectangle((x, y, x+w, y+h), fill=fill)
        return shape

    def text(self, x, y, w, h, text, size=30, color=None, bold=False):
        color = color or self.fg
        lines = wrap(text, w-8, size, bold)
        line_height = size * 1.25
        if len(lines) * line_height > h + 3:
            raise ValueError(f'第{self.index+1}页文本溢出: {text[:40]} ({len(lines)*line_height:.0f}>{h})')
        box = self.slide.shapes.add_textbox(Inches(x/120), Inches(y/120), Inches(w/120), Inches(h/120))
        tf = box.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Microsoft YaHei'
            p.font.size = Pt(size*0.6)
            p.font.bold = bold
            p.font.color.rgb = rgb(color)
            p.line_spacing = Pt(line_height*0.6)
            p.space_after = Pt(0)
            # 显式指定东亚字体，避免中文回退到不一致字体。
            for run in p.runs:
                rpr = run._r.get_or_add_rPr()
                from pptx.oxml.xmlchemy import OxmlElement as PptElement
                ea = PptElement('a:ea')
                ea.set('typeface', 'Microsoft YaHei')
                rpr.append(ea)
            self.draw.text((x, y+i*line_height), line, font=font(size,bold), fill=color, stroke_width=0)
        return box

    def picture(self, path, x, y, w, h, crop=None):
        im = Image.open(path).convert('RGB')
        if crop:
            im = im.crop(crop)
        ratio = min(w/im.width, h/im.height)
        nw, nh = int(im.width*ratio), int(im.height*ratio)
        xp, yp = int(x+(w-nw)/2), int(y+(h-nh)/2)
        stream = BytesIO()
        im.save(stream, format='PNG')
        stream.seek(0)
        self.slide.shapes.add_picture(stream, Inches(xp/120), Inches(yp/120), Inches(nw/120), Inches(nh/120))
        self.image.paste(im.resize((nw,nh),Image.Resampling.LANCZOS),(xp,yp))

    def line(self, x, y, w, color=LINE):
        self.rect(x,y,w,2,color)

    def card(self, x, y, w, title, body, h=360, accent=GREEN):
        self.rect(x,y,w,h,PALE)
        self.rect(x,y,w,5,accent)
        self.text(x+26,y+25,w-52,75,title,34,accent,True)
        self.text(x+26,y+110,w-52,h-130,body,29,INK)

    def band(self, text, y=724, color=GREEN):
        self.rect(76,y,1448,62,PALE)
        self.text(96,y+10,1408,45,text,28,color,True)

    def table(self, headers, rows, widths, y=224, rowh=66, size=25):
        x=76
        for title,w in zip(headers,widths):
            self.rect(x,y,w,rowh,INK)
            self.text(x+16,y+13,w-32,rowh-18,title,size,WHITE,True)
            x+=w
        for row_i,row in enumerate(rows):
            x=76
            for col,(value,w) in enumerate(zip(row,widths)):
                self.rect(x,y+(row_i+1)*rowh,w,rowh,WHITE if row_i%2==0 else PALE)
                self.text(x+16,y+(row_i+1)*rowh+12,w-32,rowh-18,value,size,INK,col==0)
                x+=w


def build_slide(p, i):
    if i == 0:
        p.text(76,185,1450,200,'Paper2Video\n从论文方法到教学实践',76,BG,True)
        p.text(80,414,1380,76,'模型如何配置？生成结果如何评估？',40,'#CEE1CF')
        p.rect(80,545,830,115,'#2A5145')
        p.text(106,560,780,90,'15分钟  /  模型分工 → 证据评价 → 人工决策',32,BG,True)
        p.text(80,699,1320,68,'课堂目标：判断一段教学视频是否值得继续生成和使用。',30,BG)
    elif i == 1:
        p.text(76,224,420,117,'101',92,GREEN,True)
        p.text(76,351,420,112,'组论文—作者演示视频\n平均时长 6分15秒',29,INK)
        p.text(550,234,930,96,'Paper2Video：基准与任务\nPaperTalker：生成框架',34,INK,True)
        p.text(550,363,940,180,'挑战：长文档、密集图表与公式\n课件、字幕、声音、光标、人物需要对齐',31,INK)
        p.line(76,540,1448)
        for x,n,t,b in [(76,'01','课前','看短视频，带着问题学习'),(574,'02','课中','复核错题，讨论例题'),(1072,'03','课后','修订材料，比较版本')]:
            p.text(x,569,420,60,f'{n}  {t}',36,GREEN,True)
            p.text(x,643,420,95,b,28,INK)
        p.text(76,750,1430,45,'基准统计来自论文§3.2；翻转课堂流程为本次展示设计。',23,MUTED)
    elif i == 2:
        labels=[('Slide','课件生成','Beamer + 编译反馈'),('Subtitle','讲稿 / 焦点','视觉模型读课件'),('Cursor','时空定位','坐标 + 字词时间'),('Talker','声音 / 人物','配音 + 数字人')]
        for j,(en,cn,b) in enumerate(labels):
            x=76+j*366
            p.rect(x,245,345,225,PALE)
            p.text(x+22,264,302,50,en,35,GREEN,True)
            p.text(x+22,329,302,50,cn,31,INK,True)
            p.text(x+22,396,302,60,b,24,INK)
            if j<3:p.text(x+346,330,30,50,'›',36,GREEN,True)
        p.text(76,519,1430,90,'论文：候选布局渲染 → 视觉选择；按页并行生成数字人',31,INK,True)
        p.text(76,618,1430,94,'项目：源文核对 → 讲稿审核 → 双阶段问答 → 可追溯重生成',31,GREEN,True)
        p.band('网页布局搜索默认关闭；本地任务串行，讲稿API请求可并发。')
    elif i == 3:
        rows=[('文字 / 课件','GPT-4.1¹','qwen3.8-flash','原文 → 结构 / Beamer'),('视觉 / 讲稿','GPT-4.1¹','qwen3-vl-plus','课件画面 → 口播 / 焦点'),('配音','F5-TTS','CosyVoice-v3.5-plus','口播 → 声音 / 时间戳'),('识别 / 对齐','WhisperX','Qwen3-ASR / ForcedAligner','真实音轨识别 / 本地对齐'),('光标定位','UI-TARS','PDF优先；Qwen3-VL / GUI-Owl','焦点 → 画面坐标'),('数字人','Hallo2 / FantasyTalking','MuseTalk / wan2.2-s2v','可选；新任务默认关闭'),('问答评价','Gemini-2.5-Flash²','独立 qa；qwen3-vl-plus','题库、材料 → 带证据的答案')]
        p.table(['任务角色','论文设置','当前项目可用路线','负责什么'],rows,[215,330,465,438],y=222,rowh=61,size=22)
        p.text(76,737,1448,62,'¹ 论文VLM设置；² 论文VideoLLM设置。项目型号为代码预设 / 可选路线，非历史统一配置。',22,MUTED)
    elif i == 4:
        p.rect(76,233,664,447,INK)
        snippet='单角色示例 · 不是完整任务JSON\n\n"mode": "api"\n"provider": "bailian"\n"connection_id": "class-demo"\n"model": "qwen3-vl-plus"\n"parameters": {\n    "temperature": 0.2\n}'
        p.text(102,258,608,394,snippet,27,BG)
        p.text(786,241,737,92,'模式、连接、型号、参数\n一起进入任务快照',34,GREEN,True)
        p.text(786,369,737,250,'文字 / 视觉 / qa：温度、输出上限\n云配音语速：0.5–2；本地仅默认\n讲稿并发：默认3，范围1–8\n凭据失效：重新绑定后才能继续',29,INK)
        p.band('温度示例不代表最优值；未设置参数时沿用接口默认。快照不保存密钥。')
    elif i == 5:
        p.card(76,233,690,'API 路线','减轻本地生成模型加载压力\n需要连接、材料上传和费用记录\n可使用合成服务的字词时间戳',h=356)
        p.card(800,233,724,'本地路线','准备权重、显存与运行环境\nQwen3-TTS / Qwen3-ASR\nGUI-Owl定位 / MuseTalk数字人\n文字、视觉、qa需要明确兼容服务',h=356)
        p.text(76,631,1448,76,'当前成片问答仍用本地 Qwen3-ASR；“百炼预设”是混合链路。',31,GREEN,True)
        p.band('目录里可见型号 ≠ 当前适配器可执行；未实现路线在执行前拦截。')
    elif i == 6:
        for j,(n,t,b) in enumerate([('01','口播与真实音频','数字 / 公式先规范化'),('02','字词时间轴','云时间戳 / 本地对齐'),('03','字幕与光标','按词切句 / 空间定位')]):
            x=76+j*496
            p.text(x,241,445,92,n,66,GREEN,True)
            p.text(x,354,445,60,t,34,INK,True)
            p.text(x,444,445,132,b,29,INK)
        p.text(76,596,1448,99,'PDF文字定位优先；无指示时隐藏光标；字幕区与课件分离。',31,INK,True)
        p.band('数字人开启才分段：目标18秒，严格 <20秒，检查采样连续且无重复。')
    elif i == 7:
        p.text(76,222,670,123,'¥ 5.6694',84,GREEN,True)
        p.text(76,352,640,145,'第二章历史实测累计估算\n122条去重记录 · 含失败恢复与Wan短片',28,INK)
        data=[('Wan检测 + 短片',3.069, GREEN),('TTS配音',2.3732, '#63856C'),('文字 / 视觉 / 定位',.2272212,'#94AA91')]
        for j,(t,v,c) in enumerate(data):
            yy=246+j*133
            p.text(806,yy,680,45,f'{t}    ¥{v:.4f}',28,INK,True)
            p.rect(806,yy+55,650,28,PALE)
            p.rect(806,yy+55,650*v/3.1,28,c)
        p.text(76,602,1448,89,'策略：数字人默认关；自动重生成最多2次；有效缓存才复用。',31,INK,True)
        p.band('历史标准价估算，非账户账单；不含本地算力、OSS与Codex/ImageGen。',color=AMBER)
    elif i == 8:
        rows=[('Meta Similarity','内容 / 音色是否像真人材料','研究脚本保留，网页未完整集成'),('PresentArena','哪个视频更好？交换A/B顺序','建议用于版本对照，未完整集成'),('PresentQuiz','视频覆盖了多少原文知识？','已借鉴：讲稿 + 成片双阶段问答'),('IP Memory','肖像能否关联到视频与问题？','低优先级；不是隔天学习记忆')]
        p.table(['论文指标','测量目标','当前项目边界'],rows,[340,550,558],y=241,rowh=91,size=27)
        p.band('评价维度分开报告；事实、可读性、音质与学习成效仍需其他证据。')
    elif i == 9:
        for j,(t,b) in enumerate([('核对后的原文','细节题 + 理解题\n默认20道四选一'),('固定题库','保存答案与原文摘录\n人工复核题意与答案'),('仅依据产物答题','题目 + 选项 + 材料\n不给答案和原文依据')]):
            p.card(76+j*496,241,464,t,b,h=300)
        p.text(76,589,1448,78,'有效正确题：答案正确，而且材料依据通过校验。',36,GREEN,True)
        p.band('问答分数 = 有效正确题 / 总题数 × 100%；20题时，每题改变5个百分点。')
    elif i == 10:
        p.card(76,238,690,'① 讲稿评价','只读取口播内容\n检查是否覆盖固定原文题库\n通过后才进入配音合片',h=333)
        p.card(800,238,724,'② 成片评价','最终视频 → 完整音轨 → Qwen3-ASR\n每60秒识别一段；抽取4–24张画面\n评价真实音轨转写 + 实际画面',h=333)
        p.text(76,610,1448,82,'记录原文、材料、模型、题库及语言指纹；修改后旧结果失效。',30,INK,True)
        p.band('抽帧不覆盖每个瞬间；ASR与评委可能出错。画面文字真实性仍依赖评委。',color=AMBER)
    elif i == 11:
        p.text(76,228,1448,52,'讲稿 ≥80% → 配音合片；未达到则审核 / 修改 / 显式人工继续。',31,INK,True)
        for j,(v,t,b,c) in enumerate([('<60%','自动纠错重生成','最多2次；新任务保留历史',RED),('60%–<80%','人工决策','选择接受或重新生成',AMBER),('≥80%','完成任务','教学结构门禁仍需通过',GREEN)]):
            x=76+j*496
            p.rect(x,323,464,229,PALE)
            p.text(x+22,345,420,88,v,58,c,True)
            p.text(x+22,438,420,53,t,32,INK,True)
            p.text(x+22,501,420,47,b,23,INK)
        p.text(76,593,1448,79,'课堂判断：讲稿17/20，成片14/20，下一步应该怎么做？',34,GREEN,True)
        p.band('异常 ≠ 0分：连接 / 识别失败时暂停。阈值是可调工程策略，非论文标准。')
    elif i == 12:
        p.text(76,221,750,64,'模型问答准确率 / %',33,INK,True)
        for j,(label,a,b) in enumerate([('细节题',73.8,84.2),('理解题',90.8,95.1)]):
            yy=313+j*168
            p.text(76,yy,150,50,label,30,INK,True)
            for k,(v,c) in enumerate([(a,'#718578'),(b,GREEN)]):
                by=yy+k*58
                p.rect(254,by,425*v/100,36,c)
                p.text(707,by-4,132,53,f'{v:.1f}',30,c,True)
        p.text(76,668,733,47,'灰：真人视频    绿：PaperTalker',25,MUTED)
        p.rect(880,235,2,449,LINE)
        p.text(943,222,560,62,'人工偏好 / 1–5分',33,INK,True)
        p.text(943,326,560,110,'4.6  真人',67,'#7B8F80',True)
        p.text(943,455,560,110,'3.8  PaperTalker',57,GREEN,True)
        p.text(943,596,560,88,'抽取10篇论文，10名参与者\n论文表2与图6；不同于本项目实测',25,MUTED)
        p.band('问答覆盖率和观看偏好测量不同维度；不能据此宣称学生学习效果提升。')
    elif i == 13:
        p.picture(OUT/'assets/historical-subtitle.png',76,281,914,470)
        p.rect(76,236,914,42,INK)
        p.text(90,239,880,35,'2026-10-08 历史样本 · 长字幕遮挡网络图',22,WHITE,True)
        p.text(1031,233,484,101,'17页教材 → 12页课件\n视频 25分24秒',32,GREEN,True)
        p.text(1031,363,484,137,'人工辅助恢复后解码通过\n教学交付历史评审约5/10\n存在裁切、遮挡与事实问题',25,INK)
        p.text(1031,550,484,157,'新版：流程已实现\n问答16项离线回归通过\n真实问答与学生效果待测',28,GREEN,True)
        p.text(76,765,1448,35,'证据分开：历史成片暴露问题；新版本替身测试验证流程，不提供真实效果分数。',22,AMBER)
    elif i == 14:
        p.picture(OUT/'assets/simulated-evaluation.png',633,225,891,496,crop=(112,193,1314,832))
        p.text(76,246,500,104,'1  看配置\n2  查逐题证据',36,INK,True)
        p.text(76,398,500,104,'3  判断是否重做\n4  原文复核错题',36,GREEN,True)
        p.text(76,558,500,125,'同一素材与题库\n只改变一个因素\n比较分数、可读性和费用',28,INK)
        p.band('右侧85% / 70%为离线模拟截图，只展示界面，不是项目真实评价成绩。',color=AMBER)
    elif i == 15:
        p.text(76,236,1448,136,'模型配置让过程可复现\n证据评价让结果可追溯',54,GREEN,True)
        for j,(t,b) in enumerate([('01  真实跑一次','人工复核题库\n运行新版问答与听审'),('02  做版本对照','固定材料 / 题库 / 评委\n每次只改一个因素'),('03  回到学习效果','学生观看前后测\n课中错题讨论与反馈')]):
            p.card(76+j*496,439,464,t,b,h=235)
        p.band('从“生成成功”走向“教学可信”，仍需要真实课堂数据。')
    elif i == 16:
        p.table(['问题','回答要点'],[
            ('为什么要多个模型？','能力与接口不同；便于替换、追溯和定位故障。'),
            ('80%的依据是什么？','可调的工程默认，尚未通过学生学习效果标定。'),
            ('题库摘录存在就正确？','只能验证出处存在；题意与答案仍需人工复核。'),
            ('高分等于视频无错误？','题库与抽帧有限；仍需事实审核、连续观看和听审。'),
            ('IP Memory是学习记忆？','论文做肖像—短片—问题关联，不是隔天测验。'),
        ],[430,1018],y=231,rowh=83,size=27)
        p.text(76,755,1448,35,'本页为答辩备用，不计入15分钟。',23,MUTED)
    elif i == 17:
        p.text(76,229,1448,104,'论文：Paper2Video: Automatic Video Generation from Scientific Papers\nZhu / Lin / Shou · arXiv:2510.05096v2 · 2025-10-09',29,INK,True)
        p.text(76,360,700,223,'P  论文§3.3 / §4 / §5 / 附录A.1\nC / F / R  配置、默认值与角色执行\nQ / D  问答实现与流程说明',28,INK)
        p.text(851,360,663,223,'H  第二章历史实测报告\nL  P1 / P2修复与离线验收\nV  10月11日最新评价验收日志',28,INK)
        p.text(76,619,1448,79,'完整相对路径、文件SHA256与逐页引用，见配套讲稿与 sources.json。',28,GREEN,True)
        p.band('论文页码按PDF显示页码；主要实验数据见第8页表2与图6。')


def timefmt(seconds):
    return f'{seconds//60:02d}:{seconds%60:02d}'


def build():
    OUT.mkdir(parents=True,exist_ok=True)
    assets=OUT/'assets'
    assets.mkdir(exist_ok=True)
    for source,target in [
        ('jobs/_validation/nlp_chapter2_2026-10-07/media_qa/longest-subtitle.png','historical-subtitle.png'),
        ('jobs/_validation/qa-ui/desktop-result.png','simulated-evaluation.png')]:
        if not (assets/target).exists():
            shutil.copy2(ROOT/source,assets/target)
    previews=OUT/'预览'
    previews.mkdir(exist_ok=True)
    prs=Presentation()
    prs.slide_width=Inches(W/120)
    prs.slide_height=Inches(H/120)
    prs.core_properties.title=NAME
    prs.core_properties.subject='基于论文与当前项目的翻转课堂展示；重点为模型配置与评价'
    prs.core_properties.author='Paper2Video 项目展示'
    pdf=canvas.Canvas(str(OUT/(NAME+'.pdf')),pagesize=(960,540))
    pdf.setTitle(NAME)
    pdf.setAuthor('Paper2Video 项目展示')
    markdown=['# Paper2Video 翻转课堂讲稿：模型配置与评估','',
      '版本：2026-10-11。依据本地论文 arXiv:2510.05096v2 与当前代码；16页主讲约15分钟，2页答辩备用。',
      '讲稿是可直接口述的逐页稿，并嵌入PPT备注。按约260–300中文字符/分钟排练，穿插切页、指示与约30秒互动；以实际排练为准。',
      '', '## 时间安排','', '| 页码 | 时间段 | 时长 | 主题 |','|---|---|---|---|']
    elapsed=0
    for i,s in enumerate(SLIDES):
        start=elapsed
        elapsed+=s['seconds']
        s['time']=f'{timefmt(start)}–{timefmt(elapsed)}' if s['seconds'] else '备用'
        markdown.append(f'| {i+1:02d} | {s["time"]} | {s["seconds"]}秒 | {s["title"]} |')
    assert elapsed==900
    doc=Document()
    normal=doc.styles['Normal']
    normal.font.name='Microsoft YaHei'
    normal.font.size=DocPt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal.paragraph_format.line_spacing=1.4
    doc.add_heading('Paper2Video 翻转课堂讲稿',0)
    doc.add_paragraph('模型配置与评估 · 约15分钟 · 2026-10-11')
    doc.add_paragraph('16页主讲共900秒；2页答辩备用。逐页口述稿同样嵌入PPT演讲者备注。论文结果、项目历史实测和新版离线验证分别标注。')
    table=doc.add_table(rows=1,cols=3)
    table.style='Light Shading Accent 1'
    for cell,t in zip(table.rows[0].cells,['页码 / 时间','主题','秒']):cell.text=t
    for i,s in enumerate(SLIDES[:16]):
        for cell,t in zip(table.add_row().cells,[f'{i+1:02d}  {s["time"]}',s['title'],str(s['seconds'])]):cell.text=t
    for i,s in enumerate(SLIDES):
        page=Page(prs,i,s,dark=i==0)
        build_slide(page,i)
        page.image.save(previews/f'{i+1:02d}.png')
        pdf.drawImage(ImageReader(page.image),0,0,width=960,height=540)
        pdf.showPage()
        refs='\n'.join(f'{k}：{SOURCES[k][1]}' for k in s['sources'])
        page.slide.notes_slide.notes_text_frame.text=f'第{i+1}页 · {s["time"]} · {s["seconds"]}秒\n\n{s["note"]}\n\n操作提示：{s["cue"]}\n\n来源：\n{refs}'
        markdown.extend(['',f'## 第{i+1:02d}页｜{s["title"]}',f'时间：{s["time"]}；建议{s["seconds"]}秒。','',s['note'],'',f'**操作提示：**{s["cue"]}',f'**来源：**{", ".join(s["sources"])}'])
        doc.add_page_break()
        doc.add_heading(f'第{i+1:02d}页  {s["title"]}',1)
        doc.add_paragraph(f'{s["time"]} · 建议{s["seconds"]}秒')
        doc.add_paragraph(s['note'])
        doc.add_paragraph('操作提示：'+s['cue'])
        doc.add_paragraph('来源：'+', '.join(s['sources']))
    pdf.save()
    prs.save(OUT/(NAME+'.pptx'))
    extras=[
      '## 演示安排与备用口径',
      '课前：选择3–5分钟已有视频片段，让学生带着“模型分工、评价材料、低分处理”三个问题观看；课中：按本PPT展示并讨论；课后：结合错题更新材料。这里是课堂组织建议，不是已进行学生实验。',
      '第12页互动：讲稿17/20=85%，成片14/20=70%，应进入人工决策区间，不能自动完成。若无有效音轨或连接异常，应暂停复核，不能把异常算成0分。',
      '第15页互动：可选择评价模型、字幕安全区或数字人作为实验因素，但需固定原文、题库、其他模型和产物，逐项对照；更改评价模型会重新出题，若要比较不同评委，需另行建立人工冻结的外部题库实验。',
      '现场可打开已有本机网页 http://127.0.0.1:8848/ 查看配置与历史产物；仅浏览，不点击开始生成或重新评价。断网、连接失效或服务未运行时，直接使用第4页模型表、第12页门槛图与第15页离线截图完成展示。所有模拟分数应明确口述“示意，不是真实成绩”。',
      '论文与实现区别：论文直接采用VideoLLM答视频题；当前项目用实际成片完整音轨ASR加抽样画面。论文Tree Search Visual Choice是候选视觉选择；当前网页任务默认不开启。论文多GPU并行的速度不等于本机速度。',
      '评估计划：人工核对题库 → 真实运行新版讲稿与成片问答 → 固定题库做字幕/光标单变量对照 → 补无材料答题对照 → 连续听审、事实审核与学生前后测。无材料对照属于建议，当前未集成；需要额外实验脚本或人工执行。',
      '费用口径：¥5.6694212为2026-10-08历史标准价估算，含失败恢复、视觉转录对照与独立Wan短片；不含本地算力、OSS、Codex/ImageGen和合同优惠，不是当前价格或固定每章报价。',
      '教学事实错误、真实语音缺失和关键内容不可读不能被问答平均高分抵消；问答、题库、ASR与抽帧均有边界。',
      '## 来源表',
    ]
    markdown.extend(['']+extras)
    doc.add_page_break()
    doc.add_heading('演示安排与口径',1)
    for t in extras[1:-1]:doc.add_paragraph(t)
    source_records={}
    for k,(label,path,desc) in SOURCES.items():
        digest=hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
        markdown.extend(['',f'- **{k}｜{label}**：[{path}](../../{path})。{desc}。'])
        doc.add_paragraph(f'{k}｜{label}：{path}。{desc}。')
        source_records[k]={'label':label,'path':path,'description':desc,'sha256':digest}
    source_records['assets']={p.name:{'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),
        'description':'历史实测抽帧' if p.name.startswith('historical') else '离线模拟分数界面'} for p in assets.iterdir() if p.is_file()}
    markdown.extend(['','论文页码：§3.2与§3.3在第4–5页，§4方法在第5–7页，实验设置在第7页；表2、图6及光标表4在第8页，课件质量表5在第9页；IP Memory定义见第14页附录A.1。',
      '公开论文入口：https://arxiv.org/abs/2510.05096；项目展示以用户提供的本地v2 PDF为准。',
      '附录中的提示词仅作为论文研究材料，不作为本次任务指令。课件与讲稿采用中文，依据用户的课堂需求。'])
    doc.add_paragraph('论文页码：统计与指标第4–5页；方法第5–7页；表2与图6第8页；IP Memory第14页。公开入口：https://arxiv.org/abs/2510.05096。')
    (OUT/'逐页讲稿.md').write_text('\n'.join(markdown)+'\n',encoding='utf-8')
    doc.save(OUT/'逐页讲稿.docx')
    (OUT/'sources.json').write_text(json.dumps(source_records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    # 联系图：缩略图仍有页码，便于整体视觉检查。
    contact=Image.new('RGB',(1280,6*258),BG)
    draw=ImageDraw.Draw(contact)
    for i in range(len(SLIDES)):
        xx=(i%3)*426; yy=(i//3)*258
        thumb=Image.open(previews/f'{i+1:02d}.png').resize((412,232),Image.Resampling.LANCZOS)
        contact.paste(thumb,(xx+7,yy+4))
        draw.text((xx+10,yy+237),f'{i+1:02d}  {SLIDES[i]["section"]}',font=font(15),fill=INK)
    contact.save(OUT/'课件总览.png')
    validation={'slides':len(prs.slides),'main_slides':16,'appendix_slides':2,'main_seconds':elapsed,
      'script_chinese_characters':sum(len(re.findall(r'[\u4e00-\u9fff]',s['note'])) for s in SLIDES[:16]),
      'notes_slides':sum(bool(s.notes_slide.notes_text_frame.text) for s in prs.slides),
      'screenshot_score_type':'离线模拟，不是真实成绩',
      'historical_media':'2026-10-08旧版样本，不代表当前新版',
      'outputs':[p.name for p in OUT.iterdir() if p.is_file()]}
    (OUT/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(validation,ensure_ascii=False,indent=2))


if __name__=='__main__':
    build()
