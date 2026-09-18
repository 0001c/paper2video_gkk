"""smoke test：验证 paper_ingest 对各类输入的处理效果。

用法：
    python scripts/smoke_ingest.py                  # 项目内可得的样本
    python scripts/smoke_ingest.py <某.doc 路径>     # 追加测 OLE2 老格式

样本：
  .tex  → assets/papers/gkk_mengxin/main.tex（项目内，真实）
  .pdf  → docs/Paper2Video.pdf（项目内，真实，19 页英文论文）
  .docx → 现场生成（系统内无现成样本，含标题样式，模拟真实论文）
  .doc  → 需命令行传入（python-docx 读不了该格式，故单独验证）
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import paper_ingest as pi  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, ".workbuddy", "ingest_smoke")
os.makedirs(OUT, exist_ok=True)


def make_docx_sample():
    """造一个带标题样式的中文论文样本 —— 用来验证层级保留。"""
    import docx

    document = docx.Document()
    document.add_heading("镜像与幻象：黑色电影理论的视觉谱系与本土化跨国重构", 0)
    document.add_heading("摘要", 1)
    document.add_paragraph(
        "黑色电影自1946年被法国影评人尼诺·法兰克命名以来，始终展现出难以被简单归类的理论特质。"
        "它并非好莱坞制片厂自觉建构的类型标签，而是法国知识分子在战后的回溯性美学总结。")
    document.add_heading("1 引言", 1)
    document.add_paragraph(
        "自1946年首次将“黑色电影”这一术语引入电影批评话语以来，这个概念便展现出难以被轻易归类的特质。")
    document.add_heading("1.1 研究背景", 2)
    document.add_paragraph(
        "从理论发生的历史语境来看，黑色电影最初并非美国好莱坞制片厂系统自觉标榜的类型标签。")
    document.add_heading("参考文献", 1)
    document.add_paragraph("[1] 尼诺·法兰克. 黑色电影的家族树[J]. 电影评论, 1946.")
    path = os.path.join(OUT, "sample.docx")
    document.save(path)
    return path


def run(name, source):
    print()
    print("=" * 72)
    print("  %s  <-  %s" % (name.upper(), os.path.basename(source)))
    print("=" * 72)
    try:
        text, meta = pi.ingest(source, os.path.join(OUT, name))
        for key, value in meta.items():
            if key == "source":
                continue
            print("   %-16s %s" % (key, value))
        if text:
            preview = text[:300].replace("\n", "\n      ")
            print("   --- 正文前 300 字符 ---")
            print("      " + preview)
    except Exception as exc:
        print("   失败 %s: %s" % (type(exc).__name__, str(exc)[:260]))


def check_pdf_line_merge():
    """专项：确认行合并真的把断句接回来了。"""
    print()
    print("=" * 72)
    print("  专项验证：PDF 行合并")
    print("=" * 72)
    import pymupdf

    pdf_path = os.path.join(ROOT, "docs", "Paper2Video.pdf")
    if not os.path.exists(pdf_path):
        print("   样本不存在，跳过")
        return
    doc = pymupdf.open(pdf_path)
    raw_lines, merged = [], []
    for idx in range(min(4, doc.page_count)):
        raw_lines += [ln.strip() for ln in doc[idx].get_text().splitlines() if ln.strip()]
        merged += pi._merge_pdf_lines(doc[idx].get_text().splitlines())
    doc.close()
    print("   逐行拼接 : %d 行，平均每行 %.1f 字符" % (
        len(raw_lines), sum(map(len, raw_lines)) / max(len(raw_lines), 1)))
    print("   合并后   : %d 段，平均每段 %.1f 字符" % (
        len(merged), sum(map(len, merged)) / max(len(merged), 1)))
    print("   （合并后平均长度应显著大于逐行、但也不宜过大 ——")
    print("     过小=没接上，过大=过度合并把整页糊成一段）")


def main(argv):
    samples = [
        ("tex", os.path.join(ROOT, "assets", "papers", "gkk_mengxin", "main.tex")),
        ("docx", make_docx_sample()),
        ("pdf", os.path.join(ROOT, "docs", "Paper2Video.pdf")),
    ]
    if len(argv) > 1:
        samples.append(("doc", argv[1]))
    for name, source in samples:
        if not os.path.exists(source):
            print("   跳过 %s（样本不存在：%s）" % (name, source))
            continue
        run(name, source)
    check_pdf_line_merge()
    print()
    print("输出目录：", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
